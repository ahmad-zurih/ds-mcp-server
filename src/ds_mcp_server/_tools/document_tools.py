"""
Document / file-intelligence tools.

Extract text, tables and structured previews from PDFs, Word documents,
Excel workbooks and images. Every heavy dependency (pypdf, pdfplumber,
python-docx, pytesseract/Pillow) is imported lazily at call time so the base
install stays slim and a missing library yields a clear, actionable message
instead of an import crash at server start-up.
"""
from __future__ import annotations

import os

# Cap the size of text returned to the model so a huge document cannot blow up
# the context window. Individual tools may override with their own argument.
_DEFAULT_MAX_CHARS = 12000


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------


def _check_file(path: str, exts: tuple[str, ...] | None = None) -> str | None:
    """Return an error string if ``path`` is missing or the wrong type, else None."""
    if not path or not os.path.exists(path):
        return f"Error: file not found at: {path}"
    if not os.path.isfile(path):
        return f"Error: not a file: {path}"
    if exts and not path.lower().endswith(exts):
        return (
            f"Error: unsupported file type for this tool. Expected one of "
            f"{', '.join(exts)}, got: {os.path.basename(path)}"
        )
    return None


def _truncate(text: str, max_chars: int) -> str:
    """Truncate ``text`` to ``max_chars`` with a visible marker."""
    if max_chars > 0 and len(text) > max_chars:
        omitted = len(text) - max_chars
        return text[:max_chars] + f"\n\n... [truncated {omitted:,} characters]"
    return text


def _parse_page_range(pages: str | None, total: int) -> list[int]:
    """
    Convert a 1-based page spec like ``"1,3,5-8"`` into 0-based indices.

    Returns every page when ``pages`` is falsy. Out-of-range values are ignored.
    """
    if not pages or not str(pages).strip():
        return list(range(total))
    wanted: set[int] = set()
    for chunk in str(pages).split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        if "-" in chunk:
            lo_s, _, hi_s = chunk.partition("-")
            try:
                lo, hi = int(lo_s), int(hi_s)
            except ValueError:
                continue
            for p in range(lo, hi + 1):
                if 1 <= p <= total:
                    wanted.add(p - 1)
        else:
            try:
                p = int(chunk)
            except ValueError:
                continue
            if 1 <= p <= total:
                wanted.add(p - 1)
    return sorted(wanted)


def _tables_to_markdown(tables: list[list[list]], page_no: int) -> str:
    """Render pdfplumber-style extracted tables (list of rows) as markdown."""
    import pandas as pd

    out: list[str] = []
    for t_idx, rows in enumerate(tables, 1):
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        body = [[str(c) if c is not None else "" for c in r] for r in rows[1:]]
        try:
            df = pd.DataFrame(body, columns=header)
            out.append(
                f"Page {page_no}, table {t_idx} "
                f"({df.shape[0]} rows x {df.shape[1]} cols):\n"
                + df.to_markdown(index=False)
            )
        except Exception:
            # Fall back to a raw render if the table is ragged.
            lines = [" | ".join(str(c) for c in row) for row in rows]
            out.append(f"Page {page_no}, table {t_idx}:\n" + "\n".join(lines))
    return "\n\n".join(out)


# ---------------------------------------------------------------------------
# read_pdf
# ---------------------------------------------------------------------------


def read_pdf_impl(
    file_path: str,
    pages: str | None = None,
    include_tables: bool = False,
    max_chars: int = _DEFAULT_MAX_CHARS,
) -> str:
    """
    Extract text (and optionally tables) from a PDF.

    pages:          1-based selection, e.g. "1,3,5-8"; omit for all pages.
    include_tables: also pull tables out via pdfplumber and append as markdown.
    """
    err = _check_file(file_path, (".pdf",))
    if err:
        return err

    try:
        from pypdf import PdfReader
    except ImportError:
        return (
            "pypdf is not installed. Install the documents extra:\n"
            "  pip install 'ds-mcp-server[documents]'"
        )

    try:
        reader = PdfReader(file_path)
    except Exception as exc:
        return f"Error opening PDF: {exc}"

    total = len(reader.pages)
    indices = _parse_page_range(pages, total)
    if not indices:
        return f"Error: no valid pages selected (PDF has {total} page(s))."

    parts: list[str] = [f"PDF: {os.path.basename(file_path)} ({total} page(s))"]
    for i in indices:
        try:
            text = reader.pages[i].extract_text() or ""
        except Exception as exc:
            text = f"[error extracting page {i + 1}: {exc}]"
        parts.append(f"\n--- Page {i + 1} ---\n{text.strip()}")

    result = "\n".join(parts)

    if include_tables:
        table_text = _extract_pdf_tables(file_path, indices)
        if table_text:
            result += "\n\n=== Tables ===\n" + table_text

    return _truncate(result, max_chars)


def _extract_pdf_tables(file_path: str, indices: list[int]) -> str:
    """Extract tables from the given 0-based page indices using pdfplumber."""
    try:
        import pdfplumber
    except ImportError:
        return "[pdfplumber not installed; run pip install 'ds-mcp-server[documents]' for tables]"

    chunks: list[str] = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for i in indices:
                if i >= len(pdf.pages):
                    continue
                tables = pdf.pages[i].extract_tables() or []
                if tables:
                    chunks.append(_tables_to_markdown(tables, i + 1))
    except Exception as exc:
        return f"[error extracting tables: {exc}]"
    return "\n\n".join(chunks)


# ---------------------------------------------------------------------------
# extract_tables_from_pdf
# ---------------------------------------------------------------------------


def extract_tables_from_pdf_impl(
    file_path: str,
    pages: str | None = None,
    max_chars: int = _DEFAULT_MAX_CHARS,
) -> str:
    """
    Pull structured tables out of a PDF and return them as markdown tables.

    pages: 1-based selection, e.g. "1,3,5-8"; omit for all pages.
    """
    err = _check_file(file_path, (".pdf",))
    if err:
        return err

    try:
        import pdfplumber
    except ImportError:
        return (
            "pdfplumber is not installed. Install the documents extra:\n"
            "  pip install 'ds-mcp-server[documents]'"
        )

    try:
        with pdfplumber.open(file_path) as pdf:
            total = len(pdf.pages)
            indices = _parse_page_range(pages, total)
            chunks: list[str] = []
            for i in indices:
                tables = pdf.pages[i].extract_tables() or []
                if tables:
                    chunks.append(_tables_to_markdown(tables, i + 1))
    except Exception as exc:
        return f"Error extracting tables: {exc}"

    if not chunks:
        return (
            f"No tables detected in {os.path.basename(file_path)}"
            + (f" (pages {pages})" if pages else "")
            + "."
        )
    header = f"Tables from {os.path.basename(file_path)}:\n\n"
    return _truncate(header + "\n\n".join(chunks), max_chars)


# ---------------------------------------------------------------------------
# read_docx
# ---------------------------------------------------------------------------


def read_docx_impl(file_path: str, max_chars: int = _DEFAULT_MAX_CHARS) -> str:
    """Extract text (paragraphs and tables) from a Word .docx document."""
    err = _check_file(file_path, (".docx",))
    if err:
        return err

    try:
        import docx  # python-docx
    except ImportError:
        return (
            "python-docx is not installed. Install the documents extra:\n"
            "  pip install 'ds-mcp-server[documents]'"
        )

    try:
        document = docx.Document(file_path)
    except Exception as exc:
        return f"Error opening Word document: {exc}"

    lines: list[str] = [f"Word document: {os.path.basename(file_path)}"]

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    if paragraphs:
        lines.append("\n".join(paragraphs))

    for t_idx, table in enumerate(document.tables, 1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ]
        if rows:
            lines.append(f"\n[Table {t_idx}]\n" + "\n".join(rows))

    if len(lines) == 1:
        lines.append("(no extractable text found)")

    return _truncate("\n\n".join(lines), max_chars)


# ---------------------------------------------------------------------------
# read_excel_sheets
# ---------------------------------------------------------------------------


def read_excel_sheets_impl(file_path: str, preview_rows: int = 5) -> str:
    """
    List every sheet in an Excel workbook and preview the first rows of each.

    preview_rows: number of rows to show per sheet (1-50, default 5).
    """
    err = _check_file(file_path, (".xlsx", ".xls"))
    if err:
        return err

    try:
        import pandas as pd
    except ImportError:  # pandas is a base dep, but stay defensive.
        return "pandas is not installed."

    preview_rows = max(1, min(int(preview_rows), 50))

    try:
        xls = pd.ExcelFile(file_path)
    except ImportError:
        return (
            "openpyxl is not installed. Install the documents extra:\n"
            "  pip install 'ds-mcp-server[documents]'"
        )
    except Exception as exc:
        return f"Error opening Excel workbook: {exc}"

    sheet_names = xls.sheet_names
    parts: list[str] = [
        f"Excel workbook: {os.path.basename(file_path)} "
        f"({len(sheet_names)} sheet(s): {', '.join(map(str, sheet_names))})"
    ]
    for name in sheet_names:
        try:
            df = xls.parse(name, nrows=preview_rows)
            parts.append(
                f"\n--- Sheet '{name}' — {df.shape[1]} column(s) ---\n"
                f"Columns: {', '.join(map(str, df.columns))}\n"
                + df.to_markdown(index=False)
            )
        except Exception as exc:
            parts.append(f"\n--- Sheet '{name}' ---\n[error reading sheet: {exc}]")

    return "\n".join(parts)


# ---------------------------------------------------------------------------
# ocr_image
# ---------------------------------------------------------------------------


def ocr_image_impl(file_path: str, lang: str = "eng", max_chars: int = _DEFAULT_MAX_CHARS) -> str:
    """
    Run OCR on an image (screenshot, scan or photo) and return the text.

    lang: tesseract language code(s), e.g. 'eng' or 'eng+deu'.
    Requires the `tesseract` binary plus:  pip install 'ds-mcp-server[ocr]'
    """
    err = _check_file(file_path, (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".gif", ".webp"))
    if err:
        return err

    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return (
            "OCR dependencies are not installed. Install the ocr extra:\n"
            "  pip install 'ds-mcp-server[ocr]'\n"
            "You also need the system tesseract binary "
            "(e.g. `apt install tesseract-ocr` or `brew install tesseract`)."
        )

    try:
        with Image.open(file_path) as img:
            text = pytesseract.image_to_string(img, lang=lang)
    except pytesseract.TesseractNotFoundError:
        return (
            "The tesseract binary was not found on PATH. Install it with your "
            "package manager, e.g. `apt install tesseract-ocr` or "
            "`brew install tesseract`."
        )
    except Exception as exc:
        return f"OCR error: {exc}"

    text = text.strip()
    if not text:
        return f"OCR produced no text from {os.path.basename(file_path)}."
    header = f"OCR text from {os.path.basename(file_path)} (lang={lang}):\n\n"
    return _truncate(header + text, max_chars)


# ---------------------------------------------------------------------------
# summarize_document
# ---------------------------------------------------------------------------


def _extract_document_text(file_path: str) -> tuple[bool, str]:
    """Extract plain text from a supported document type (pdf/docx/txt/md/csv).

    Returns ``(ok, text)``. When ``ok`` is False, ``text`` is an error or
    install-hint message that should be surfaced to the caller verbatim rather
    than treated as document content.
    """
    lower = file_path.lower()
    if lower.endswith(".pdf"):
        out = read_pdf_impl(file_path, max_chars=0)
        ok = not out.startswith(("Error", "pypdf is not installed"))
        return ok, out
    if lower.endswith(".docx"):
        out = read_docx_impl(file_path, max_chars=0)
        ok = not out.startswith(("Error", "python-docx is not installed"))
        return ok, out
    if lower.endswith((".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv")):
        for enc in ("utf-8", "latin1"):
            try:
                with open(file_path, "r", encoding=enc) as fh:
                    return True, fh.read()
            except UnicodeDecodeError:
                continue
        return False, "Error: could not decode text file."
    return False, (
        f"Error: unsupported document type for summarization: "
        f"{os.path.basename(file_path)}. Supported: pdf, docx, txt, md, csv."
    )


def summarize_document_impl(
    file_path: str,
    max_chars: int = _DEFAULT_MAX_CHARS,
    chunk_chars: int = 4000,
) -> str:
    """
    Prepare a (possibly long) document for LLM summarization.

    Extracts the full text, splits it into labelled chunks and returns as much
    as fits in ``max_chars`` together with an instruction for the model to
    write the summary. The calling LLM performs the actual summarization.
    """
    err = _check_file(file_path)
    if err:
        return err

    ok, text = _extract_document_text(file_path)
    if not ok:
        return text

    text = text.strip()
    if not text:
        return f"No extractable text found in {os.path.basename(file_path)}."

    chunk_chars = max(500, int(chunk_chars))
    chunks = [text[i : i + chunk_chars] for i in range(0, len(text), chunk_chars)]

    header = (
        f"Document '{os.path.basename(file_path)}' — {len(text):,} characters, "
        f"{len(chunks)} chunk(s).\n"
        "Please write a concise summary of the content below. If the text is "
        "truncated, summarize what is present and note that it was cut off.\n"
    )

    body_parts: list[str] = []
    used = len(header)
    for idx, chunk in enumerate(chunks, 1):
        label = f"\n--- Chunk {idx}/{len(chunks)} ---\n"
        if used + len(label) + len(chunk) > max_chars > 0:
            remaining = max_chars - used - len(label)
            if remaining > 200:
                body_parts.append(label + chunk[:remaining])
            body_parts.append(
                f"\n\n... [document truncated; {len(chunks) - idx + 1} "
                "chunk(s) omitted]"
            )
            break
        body_parts.append(label + chunk)
        used += len(label) + len(chunk)

    return header + "".join(body_parts)
